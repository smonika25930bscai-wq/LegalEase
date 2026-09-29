import os
import time
import base64
import io
import re

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel

from google import genai
from google.genai import types

from dotenv import load_dotenv

from docx import Document
from docx.shared import Inches, Pt

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Image,
)
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics


# ============================================================
# Environment
# ============================================================

load_dotenv()


# ============================================================
# Router
# ============================================================

router = APIRouter()


# ============================================================
# Gemini Configuration
# ============================================================

GEMINI_API_KEY = os.getenv(
    "GEMINI_API_KEY",
    ""
).strip()

GEMINI_MODEL = os.getenv(
    "GEMINI_MODEL",
    "gemini-2.5-flash"
).strip()


# ============================================================
# Request Models
# ============================================================

class GenerateRequest(BaseModel):
    prompt: str


class ExportRequest(BaseModel):
    text: str
    document_type: str = "Legal Document"
    logo_base64: str | None = None


# ============================================================
# Home
# ============================================================

@router.get("/")
def home():
    return {
        "message": "LegalEase Backend is running",
        "model": GEMINI_MODEL
    }


# ============================================================
# Health Check
# ============================================================

@router.get("/health")
def health():
    return {
        "status": "ok",
        "gemini_configured": bool(GEMINI_API_KEY),
        "model": GEMINI_MODEL
    }


# ============================================================
# Helper: Temporary Gemini errors
# ============================================================

def is_temporary_gemini_error(error: Exception) -> bool:

    error_text = str(error).lower()

    temporary_errors = [
        "503",
        "unavailable",
        "service unavailable",
        "high demand",
        "overloaded",
        "temporarily unavailable",
        "resource exhausted",
        "429",
        "too many requests",
    ]

    return any(
        message in error_text
        for message in temporary_errors
    )


# ============================================================
# Generate Document
# ============================================================

@router.post("/generate")
def generate(request: GenerateRequest):

    if not GEMINI_API_KEY:
        raise HTTPException(
            status_code=500,
            detail=(
                "GEMINI_API_KEY is not configured. "
                "Please add GEMINI_API_KEY to .env."
            )
        )

    if not request.prompt.strip():
        raise HTTPException(
            status_code=400,
            detail="Prompt cannot be empty."
        )

    try:
        client = genai.Client(
            api_key=GEMINI_API_KEY
        )

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Could not initialize Gemini client: {str(e)}"
        )

    max_attempts = 3
    last_error = None

    for attempt in range(1, max_attempts + 1):

        try:

            response = client.models.generate_content(

                model=GEMINI_MODEL,

                contents=request.prompt,

                config=types.GenerateContentConfig(

                    temperature=0.25,

                    max_output_tokens=8192,

                    candidate_count=1,
                ),
            )

            try:
                generated_text = response.text
            except Exception:
                generated_text = None

            if not generated_text:
                raise HTTPException(
                    status_code=502,
                    detail=(
                        "Gemini returned an empty response. "
                        "Please try again."
                    )
                )

            generated_text = generated_text.strip()

            if not generated_text:
                raise HTTPException(
                    status_code=502,
                    detail=(
                        "Gemini returned an empty document. "
                        "Please try again."
                    )
                )

            return {
                "success": True,
                "content": generated_text,
                "model": GEMINI_MODEL,
            }

        except HTTPException:
            raise

        except Exception as e:

            last_error = e

            if is_temporary_gemini_error(e):

                if attempt < max_attempts:

                    wait_seconds = attempt * 3

                    time.sleep(wait_seconds)

                    continue

                raise HTTPException(
                    status_code=503,
                    detail=(
                        "Gemini is temporarily unavailable "
                        "because the model is experiencing "
                        "high demand. "
                        "The backend tried multiple times. "
                        "Please wait a little and try again."
                    )
                )

            raise HTTPException(
                status_code=500,
                detail=f"Gemini API error: {str(e)}"
            )

    raise HTTPException(
        status_code=500,
        detail=(
            "Document generation failed. "
            f"Gemini error: {str(last_error)}"
        )
    )


# ============================================================
# Helper: Validate Export Text
# ============================================================

def validate_export_text(text: str):

    if not text or not text.strip():
        raise HTTPException(
            status_code=400,
            detail="Document text cannot be empty."
        )


# ============================================================
# Helper: Decode Logo
# ============================================================

def decode_logo(logo_base64: str | None):

    if not logo_base64:
        return None

    try:

        if "," in logo_base64:
            logo_base64 = logo_base64.split(",", 1)[1]

        return base64.b64decode(logo_base64)

    except Exception:

        return None


# ============================================================
# Helper: Clean Filename
# ============================================================

def clean_filename(name: str) -> str:

    name = re.sub(
        r"[^a-zA-Z0-9_-]+",
        "_",
        name
    )

    name = name.strip("_")

    if not name:
        name = "LegalEase_Document"

    return name


# ============================================================
# EXPORT TXT
# ============================================================

@router.post("/export/txt")
def export_txt(request: ExportRequest):

    validate_export_text(request.text)

    filename = clean_filename(
        request.document_type
    )

    content = request.text.encode(
        "utf-8"
    )

    return Response(
        content=content,
        media_type="text/plain; charset=utf-8",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}.txt"'
            )
        }
    )


# ============================================================
# EXPORT DOCX
# ============================================================

@router.post("/export/docx")
def export_docx(request: ExportRequest):

    validate_export_text(request.text)

    try:

        document = Document()

        # ----------------------------------------------------
        # Page margins
        # ----------------------------------------------------

        section = document.sections[0]

        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        # ----------------------------------------------------
        # Logo
        # ----------------------------------------------------

        logo_bytes = decode_logo(
            request.logo_base64
        )

        if logo_bytes:

            try:

                logo_stream = io.BytesIO(
                    logo_bytes
                )

                paragraph = document.add_paragraph()

                paragraph.alignment = 1

                run = paragraph.add_run()

                run.add_picture(
                    logo_stream,
                    width=Inches(1.4)
                )

            except Exception:
                pass

        # ----------------------------------------------------
        # Document title
        # ----------------------------------------------------

        title = document.add_paragraph()

        title.alignment = 1

        title_run = title.add_run(
            request.document_type.upper()
        )

        title_run.bold = True
        title_run.font.size = Pt(16)

        document.add_paragraph()

        # ----------------------------------------------------
        # Document content
        # ----------------------------------------------------

        for line in request.text.splitlines():

            line = line.strip()

            if not line:
                document.add_paragraph()
                continue

            paragraph = document.add_paragraph()

            # Numbered section
            if re.match(
                r"^\d+\.\s+",
                line
            ):

                run = paragraph.add_run(
                    line
                )

                run.bold = True
                run.font.size = Pt(12)

            # Heading
            elif (
                line.isupper()
                and len(line) < 120
            ):

                run = paragraph.add_run(
                    line
                )

                run.bold = True
                run.font.size = Pt(13)

            # Bullet
            elif line.startswith("- "):

                paragraph.style = (
                    document.styles["List Bullet"]
                )

                run = paragraph.add_run(
                    line[2:]
                )

                run.font.size = Pt(11)

            else:

                run = paragraph.add_run(
                    line
                )

                run.font.size = Pt(11)

        # ----------------------------------------------------
        # Save to memory
        # ----------------------------------------------------

        output = io.BytesIO()

        document.save(output)

        output.seek(0)

        filename = clean_filename(
            request.document_type
        )

        return Response(

            content=output.getvalue(),

            media_type=(
                "application/vnd.openxmlformats-"
                "officedocument.wordprocessingml.document"
            ),

            headers={
                "Content-Disposition": (
                    f'attachment; filename="{filename}.docx"'
                )
            }
        )

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=f"DOCX export failed: {str(e)}"
        )


# ============================================================
# EXPORT PDF
# ============================================================

@router.post("/export/pdf")
def export_pdf(request: ExportRequest):

    validate_export_text(request.text)

    try:

        output = io.BytesIO()

        # ----------------------------------------------------
        # PDF document
        # ----------------------------------------------------

        pdf = SimpleDocTemplate(

            output,

            pagesize=A4,

            rightMargin=50,

            leftMargin=50,

            topMargin=50,

            bottomMargin=50,
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(

            "LegalTitle",

            parent=styles["Title"],

            alignment=TA_CENTER,

            fontSize=16,

            leading=20,

            spaceAfter=15,
        )

        heading_style = ParagraphStyle(

            "LegalHeading",

            parent=styles["Heading2"],

            fontSize=13,

            leading=16,

            spaceBefore=10,

            spaceAfter=6,
        )

        paragraph_style = ParagraphStyle(

            "LegalParagraph",

            parent=styles["BodyText"],

            fontSize=10.5,

            leading=15,

            spaceAfter=7,
        )

        bullet_style = ParagraphStyle(

            "LegalBullet",

            parent=styles["BodyText"],

            fontSize=10.5,

            leading=15,

            leftIndent=15,

            firstLineIndent=-8,

            spaceAfter=5,
        )

        story = []

        # ----------------------------------------------------
        # Logo
        # ----------------------------------------------------

        logo_bytes = decode_logo(
            request.logo_base64
        )

        if logo_bytes:

            try:

                logo_stream = io.BytesIO(
                    logo_bytes
                )

                logo = Image(
                    logo_stream,
                    width=1.2 * inch,
                    height=1.2 * inch,
                )

                logo.hAlign = "CENTER"

                story.append(logo)

                story.append(
                    Spacer(1, 10)
                )

            except Exception:
                pass

        # ----------------------------------------------------
        # Title
        # ----------------------------------------------------

        story.append(
            Paragraph(
                request.document_type.upper(),
                title_style
            )
        )

        # ----------------------------------------------------
        # Content
        # ----------------------------------------------------

        for line in request.text.splitlines():

            line = line.strip()

            if not line:

                story.append(
                    Spacer(1, 6)
                )

                continue

            safe_line = (
                line
                .replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )

            # Numbered section
            if re.match(
                r"^\d+\.\s+",
                line
            ):

                story.append(
                    Paragraph(
                        f"<b>{safe_line}</b>",
                        heading_style
                    )
                )

            # Uppercase heading
            elif (
                line.isupper()
                and len(line) < 120
            ):

                story.append(
                    Paragraph(
                        f"<b>{safe_line}</b>",
                        heading_style
                    )
                )

            # Bullet
            elif line.startswith("- "):

                bullet_text = safe_line[2:]

                story.append(
                    Paragraph(
                        f"• {bullet_text}",
                        bullet_style
                    )
                )

            else:

                story.append(
                    Paragraph(
                        safe_line,
                        paragraph_style
                    )
                )

        # ----------------------------------------------------
        # Build PDF
        # ----------------------------------------------------

        pdf.build(story)

        output.seek(0)

        filename = clean_filename(
            request.document_type
        )

        return Response(

            content=output.getvalue(),

            media_type="application/pdf",

            headers={
                "Content-Disposition": (
                    f'attachment; filename="{filename}.pdf"'
                )
            }
        )

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=f"PDF export failed: {str(e)}"
        )